#!/usr/bin/env python3
"""Converts docs/bao-cao.md (single source of truth) into report/bao-cao.docx.

Tailored to the known structure of this one report (cover section, TOC
section, chapter headings, tables, code blocks) rather than a generic
Markdown engine. Re-run via report/build.sh whenever docs/bao-cao.md changes.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "docs" / "bao-cao.md"
OUT_PATH = ROOT / "report" / "bao-cao.docx"

SERIF = "Times New Roman"
MONO = "Courier New"

PLACEHOLDER_RE = re.compile(r"(SV X[AÁ]C NH[AẬ]N|SV ĐI[ÈỀ]N|CH[ƯUƯ]A L[ÀA]M)")
INLINE_RE = re.compile(
    r"(?P<bold>\*\*(?:[^*]|\*(?!\*))+?\*\*)"
    r"|(?P<code>`[^`]+?`)"
    r"|(?P<italic>\*[^*]+?\*)"
    r"|(?P<url>https?://[^\s\]\)”\"]+)"
)

USABLE_WIDTH_CM = 21.0 - 3.0 - 2.5  # A4 width minus left/right margins

TABLE_WIDTH_RATIOS = {
    ("Bước", "Dòng mã giả", "Giải thích hiển thị trên Player", "Mảng màu", "Cây"): [
        0.07,
        0.12,
        0.42,
        0.24,
        0.15,
    ],
    ("Thuật toán", "Số bước", "Phép toán", "Kết quả", "Checker"): [
        0.14,
        0.11,
        0.11,
        0.42,
        0.22,
    ],
    ("Thuật toán", "Số màu dùng", "Số bước", "Phép toán", "Kết quả"): [
        0.17,
        0.16,
        0.14,
        0.14,
        0.39,
    ],
    ("Thuật toán", "Kết luận chắc chắn bài k màu?", "Độ phức tạp lý thuyết / trong app", "Giới hạn n trong app"): [
        0.14,
        0.22,
        0.32,
        0.32,
    ],
    ("Nhóm", "File", "n / m", "Chế độ, k", "Kỳ vọng", "Mục đích"): [
        0.11,
        0.19,
        0.10,
        0.14,
        0.11,
        0.35,
    ],
}


def strip_url_trailing_punct(url: str) -> tuple[str, str]:
    trail = ""
    while url and url[-1] in ".,;:)»":
        trail = url[-1] + trail
        url = url[:-1]
    return url, trail


def add_hyperlink(paragraph, url: str, text: str, size_pt: float):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), SERIF)
    rFonts.set(qn("w:hAnsi"), SERIF)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def emit_plain(paragraph, text: str, bold=False, italic=False, code=False, size_pt=None, color=None):
    """Add `text` as one or more runs, splitting out yellow-highlighted placeholders."""
    parts = PLACEHOLDER_RE.split(text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = MONO
            run.font.size = Pt(size_pt - 1.5 if size_pt else 11)
        elif size_pt:
            run.font.size = Pt(size_pt)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if PLACEHOLDER_RE.fullmatch(part):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True


def render_inline(paragraph, text: str, base_size=None, base_bold=False):
    """Parse **bold**, `code`, *italic*, and bare URLs, then add runs to `paragraph`."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            emit_plain(paragraph, text[pos : m.start()], bold=base_bold, size_pt=base_size)
        if m.group("bold"):
            emit_plain(paragraph, m.group("bold")[2:-2], bold=True, size_pt=base_size)
        elif m.group("code"):
            emit_plain(paragraph, m.group("code")[1:-1], code=True, bold=base_bold, size_pt=base_size)
        elif m.group("italic"):
            emit_plain(paragraph, m.group("italic")[1:-1], italic=True, bold=base_bold, size_pt=base_size)
        elif m.group("url"):
            url, trail = strip_url_trailing_punct(m.group("url"))
            add_hyperlink(paragraph, url, url, base_size or 13)
            if trail:
                emit_plain(paragraph, trail, bold=base_bold, size_pt=base_size)
        pos = m.end()
    if pos < len(text):
        emit_plain(paragraph, text[pos:], bold=base_bold, size_pt=base_size)


def shade_paragraph(paragraph, hex_color: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_width(cell, cm_value):
    cell.width = Cm(cm_value)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:type"), "dxa")
    tcW.set(qn("w:w"), str(int(cm_value * 567)))


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    fld_begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Nhấn Ctrl+A rồi F9 (hoặc Tham chiếu > Cập nhật bảng) để cập nhật Mục lục."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(placeholder)
    r.append(fld_end)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(placeholder)
    r.append(fld_end)


class Builder:
    def __init__(self, doc: Document):
        self.doc = doc

    # -- paragraph helpers -------------------------------------------------
    def body_paragraph(self, text, *, align=None, space_before=None, space_after=None):
        p = self.doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if space_before is not None:
            p.paragraph_format.space_before = Pt(space_before)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        render_inline(p, text)
        return p

    def heading(self, level, text):
        style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
        p = self.doc.add_paragraph(style=style)
        if level == 1:
            p.paragraph_format.page_break_before = True
        render_inline(p, text, base_size={1: 16, 2: 14, 3: 13}[level], base_bold=True)
        return p

    def bold_lead(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        render_inline(p, text, base_size=13, base_bold=True)
        return p

    def bullet(self, text, level=0):
        style = "List Bullet" if level == 0 else "List Bullet 2"
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        render_inline(p, text)
        return p

    def numbered_line(self, text, level=0):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9 + 0.6 * level)
        p.paragraph_format.space_after = Pt(2)
        render_inline(p, text)
        return p

    def quote_line(self, text, bullet=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(2)
        pPr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "12")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), "999999")
        border.append(left)
        pPr.append(border)
        shade_paragraph(p, "F2F2F2")
        prefix = "• " if bullet else ""
        if prefix:
            run = p.add_run(prefix)
            run.italic = True
            run.font.size = Pt(12)
        render_inline(p, text, base_size=12)
        for run in p.runs:
            if not run.italic:
                run.italic = True
        return p

    def image_callout(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        shade_paragraph(p, "ECECEC")
        emit_plain(p, f"[Chèn ảnh: {text}]", italic=True, size_pt=11, color="555555")
        return p

    def code_block(self, lines):
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0) if i else Pt(6)
            p.paragraph_format.line_spacing = 1.0
            shade_paragraph(p, "F2F2F2")
            leading = len(line) - len(line.lstrip(" "))
            content = " " * leading + line[leading:] if leading else line
            run = p.add_run(content if content else " ")
            run.font.name = MONO
            run.font.size = Pt(10.5)
        if lines:
            self.doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    def table(self, header, aligns, rows):
        n_cols = len(header)
        ratios = TABLE_WIDTH_RATIOS.get(tuple(header))
        if ratios is None:
            ratios = [1.0 / n_cols] * n_cols if n_cols != 2 else [0.38, 0.62]
        widths = [USABLE_WIDTH_CM * r for r in ratios]
        font_size = 10 if (n_cols >= 5 or len(rows) >= 15) else 11

        tbl = self.doc.add_table(rows=1 + len(rows), cols=n_cols)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        header_row = tbl.rows[0]
        trPr = header_row._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)
        for c, text in enumerate(header):
            cell = header_row.cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            render_inline(p, text, base_size=font_size, base_bold=True)
            shade_cell(cell, "D9D9D9")
            set_cell_width(cell, widths[c])

        align_map = {
            "l": WD_ALIGN_PARAGRAPH.LEFT,
            "c": WD_ALIGN_PARAGRAPH.CENTER,
            "r": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        for r, row in enumerate(rows, start=1):
            tbl_row = tbl.rows[r]
            for c, text in enumerate(row):
                cell = tbl_row.cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = align_map.get(aligns[c], WD_ALIGN_PARAGRAPH.LEFT)
                render_inline(p, text, base_size=font_size)
                set_cell_width(cell, widths[c])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)


# -- markdown line classification -------------------------------------------

HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)-\s+(.*)$")
NUMBERED_RE = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
IMAGE_RE = re.compile(r"^\[(?:ẢNH|ảnh)[:\.]?\s*(.*)\]$")
BOLD_ONLY_RE = re.compile(r"^\*\*([^*]+)\*\*$")
TABLE_ROW_RE = re.compile(r"^\|(.*)\|\s*$")


def split_row(line):
    inner = line.strip()[1:-1]
    return [c.strip() for c in inner.split("|")]


def parse_align(cell):
    cell = cell.strip()
    left = cell.startswith(":")
    right = cell.endswith(":")
    if left and right:
        return "c"
    if right:
        return "r"
    return "l"


def convert():
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), SERIF)

    for level, (before, after, size) in {1: (18, 10, 16), 2: (12, 6, 14), 3: (10, 4, 13)}.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = SERIF
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # footer page numbers (suppressed on the cover page only)
    section.different_first_page_header_footer = True
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.add_run("– ").font.size = Pt(10)
    add_page_number_field(footer_p)
    footer_p.add_run(" –").font.size = Pt(10)
    for run in footer_p.runs:
        run.font.name = SERIF
        run.font.size = Pt(10)
    section.first_page_footer.is_linked_to_previous = False
    for p in list(section.first_page_footer.paragraphs):
        for r in list(p.runs):
            r.text = ""

    b = Builder(doc)

    front_matter_lines = []
    i = 0
    n = len(lines)
    mode = "frontmatter"  # frontmatter -> cover -> toc -> body
    cover_lines = []

    def flush_front_matter_note():
        if not front_matter_lines:
            return
        b.heading(2, "Ghi chú biên tập (không thuộc nội dung nộp bài)")
        for raw in front_matter_lines:
            m = QUOTE_RE.match(raw)
            text = m.group(1) if m else raw
            bm = BULLET_RE.match(text)
            if bm:
                b.quote_line(bm.group(2), bullet=True)
            elif text.strip():
                b.quote_line(text)

    while i < n:
        raw = lines[i]
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        hm = HEADING_RE.match(line)
        if hm:
            level = len(hm.group(1))
            text = hm.group(2).strip()
            if level == 1:
                i += 1
                continue
            # close out whatever special section we were in before this heading
            if mode == "cover":
                render_cover(doc, b, cover_lines)
                doc.add_page_break()
                mode = "frontmatter"
            if text == "Trang bìa":
                mode = "cover"
                cover_lines = []
                i += 1
                continue
            if text == "Mục lục":
                mode = "toc"
                i += 1
                continue
            if mode == "toc":
                render_toc(doc, b)
                flush_front_matter_note()
                doc.add_page_break()
            mode = "body"
            b.heading(level - 1, text)
            i += 1
            continue

        if mode == "cover":
            cover_lines.append(line)
            i += 1
            continue
        if mode == "toc":
            # ignore the "(Word tự sinh từ các heading.)" placeholder line
            i += 1
            continue
        if mode == "frontmatter":
            front_matter_lines.append(line)
            i += 1
            continue

        # body mode -------------------------------------------------------
        if line.startswith("```"):
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            b.code_block(block)
            continue

        if TABLE_ROW_RE.match(line):
            header = split_row(line)
            i += 1
            sep = split_row(lines[i])
            aligns = [parse_align(c) for c in sep]
            i += 1
            rows = []
            while i < n and TABLE_ROW_RE.match(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            b.table(header, aligns, rows)
            continue

        if QUOTE_RE.match(line):
            qlines = []
            while i < n and QUOTE_RE.match(lines[i]):
                qm = QUOTE_RE.match(lines[i])
                qlines.append(qm.group(1))
                i += 1
            for qtext in qlines:
                bm = BULLET_RE.match(qtext)
                if bm:
                    b.quote_line(bm.group(2), bullet=True)
                elif qtext.strip():
                    b.quote_line(qtext)
            continue

        im = IMAGE_RE.match(line.strip())
        if im:
            b.image_callout(im.group(1))
            i += 1
            continue

        bm = BOLD_ONLY_RE.match(line.strip())
        if bm:
            b.bold_lead(bm.group(1))
            i += 1
            continue

        lm = BULLET_RE.match(line)
        if lm:
            indent = len(lm.group(1))
            b.bullet(lm.group(2), level=1 if indent >= 2 else 0)
            i += 1
            continue

        nm = NUMBERED_RE.match(line)
        if nm:
            b.numbered_line(f"{nm.group(2)}. {nm.group(3)}")
            i += 1
            continue

        b.body_paragraph(line)
        i += 1

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


def render_cover(doc, b, cover_lines):
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    sizes = [13, 13, 18, 20, 14, 14, 13, 13]
    space_before = [0, 0, 36, 18, 48, 6, 36, 12]
    for idx, raw in enumerate(cover_lines):
        bm = BULLET_RE.match(raw)
        text = bm.group(2) if bm else raw.strip()
        if not text:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = sizes[idx] if idx < len(sizes) else 13
        before = space_before[idx] if idx < len(space_before) else 6
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(4)
        is_title_line = bool(BOLD_ONLY_RE.match(text))
        inner = BOLD_ONLY_RE.match(text).group(1) if is_title_line else text
        render_inline(p, inner, base_size=size, base_bold=is_title_line)


def render_toc(doc, b):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("MỤC LỤC")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = SERIF
    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.space_after = Pt(6)
    add_toc_field(toc_p)
    for run in toc_p.runs:
        run.font.name = SERIF
        run.font.size = Pt(12)
        run.italic = True


if __name__ == "__main__":
    convert()
